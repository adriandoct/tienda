import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def generate_uml_docx():
    doc = docx.Document()

    # Configuración de márgenes
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Paleta de colores corporativos
    COLOR_PRIMARY = RGBColor(27, 54, 93)       # #1B365D Azul marino profundo
    COLOR_SECONDARY = RGBColor(43, 108, 176)   # #2B6CB0 Azul corporativo
    COLOR_TEAL = RGBColor(15, 118, 110)        # #0F766E Verde azulado para embebidos
    COLOR_AMBER = RGBColor(180, 83, 9)         # #B45309 Ámbar para inventario
    COLOR_TEXT = RGBColor(45, 55, 72)          # #2D3748 Gris grafito de texto
    COLOR_MUTED = RGBColor(100, 116, 139)      # #64748B Gris secundario
    COLOR_DARK = RGBColor(15, 23, 42)          # #0F172A Negro azulado
    HEX_PRIMARY = "1B365D"
    HEX_SECONDARY = "2B6CB0"
    HEX_TEAL = "0F766E"
    HEX_AMBER = "B45309"
    HEX_BG_HEADER = "1E3A8A"
    HEX_BG_SUBHEADER = "2B6CB0"
    HEX_BG_ALT = "F8FAFC"
    HEX_BG_CODE = "F1F5F9"
    HEX_BORDER = "CBD5E1"

    # Helpers XML
    def set_cell_background(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_cell_borders(cell, top=HEX_BORDER, bottom=HEX_BORDER, left=None, right=None, sz=4):
        tcPr = cell._tc.get_or_add_tcPr()
        borders_xml = f'<w:tcBorders {nsdecls("w")}>'
        borders_xml += f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>' if top else '<w:top w:val="none"/>'
        borders_xml += f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>' if bottom else '<w:bottom w:val="none"/>'
        borders_xml += f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left}"/>' if left else '<w:left w:val="none"/>'
        borders_xml += f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>' if right else '<w:right w:val="none"/>'
        borders_xml += '</w:tcBorders>'
        tcPr.append(parse_xml(borders_xml))

    def add_header_h1(text, icon="🔷"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run_icon = p.add_run(f"{icon} ")
        run_icon.font.size = Pt(15)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY

        # Línea divisoria elegante
        p_line = doc.add_paragraph()
        p_line.paragraph_format.space_before = Pt(0)
        p_line.paragraph_format.space_after = Pt(8)
        run_line = p_line.add_run("―" * 58)
        run_line.font.size = Pt(8)
        run_line.font.color.rgb = RGBColor(203, 213, 225)

    def add_header_h2(text, badge=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY

        if badge:
            run_badge = p.add_run(f"  [{badge}]")
            run_badge.font.name = "Calibri"
            run_badge.font.size = Pt(9.5)
            run_badge.font.bold = True
            if "Embebido" in badge:
                run_badge.font.color.rgb = COLOR_TEAL
            elif "Desacoplado" in badge:
                run_badge.font.color.rgb = COLOR_AMBER
            else:
                run_badge.font.color.rgb = COLOR_SECONDARY

    def add_paragraph(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_pre = p.add_run(bold_prefix)
            run_pre.font.name = "Calibri"
            run_pre.font.size = Pt(10.5)
            run_pre.font.bold = True
            run_pre.font.color.rgb = COLOR_DARK
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_TEXT
        return p

    def add_bullet(bold_label, desc, icon="• "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        run_icon = p.add_run(icon)
        run_icon.font.name = "Calibri"
        run_icon.font.size = Pt(10)
        run_icon.font.color.rgb = COLOR_SECONDARY
        run_lbl = p.add_run(bold_label)
        run_lbl.font.name = "Calibri"
        run_lbl.font.size = Pt(10.5)
        run_lbl.font.bold = True
        run_lbl.font.color.rgb = COLOR_DARK
        run_desc = p.add_run(desc)
        run_desc.font.name = "Calibri"
        run_desc.font.size = Pt(10.5)
        run_desc.font.color.rgb = COLOR_TEXT

    def add_callout(text, icon="💡", title="NOTA ARQUITECTÓNICA", color_hex="2B6CB0", bg_hex="EFF6FF"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.7)
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=140, bottom=140, left=220, right=200)

        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{color_hex}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15

        run_title = p.add_run(f"{icon} {title}\n")
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(9.5)
        run_title.font.bold = True
        run_title.font.color.rgb = COLOR_SECONDARY

        run_text = p.add_run(text)
        run_text.font.name = "Calibri"
        run_text.font.size = Pt(9.5)
        run_text.font.color.rgb = COLOR_TEXT

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(6)

    def add_json_box(code_str, caption=None):
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(4)
            p_cap.paragraph_format.space_after = Pt(2)
            run_cap = p_cap.add_run(f"📄 {caption}")
            run_cap.font.name = "Calibri"
            run_cap.font.size = Pt(9)
            run_cap.font.bold = True
            run_cap.font.color.rgb = COLOR_MUTED

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.7)
        set_cell_background(cell, HEX_BG_CODE)
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'<w:left w:val="single" w:sz="16" w:space="0" w:color="{HEX_SECONDARY}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05

        run = p.add_run(code_str)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 41, 59)

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(6)

    def add_table_fields(headers, data, widths, bg_header=HEX_BG_HEADER):
        tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Fila de cabecera
        hdr_row = tbl.rows[0]
        trPr = hdr_row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for i, h_text in enumerate(headers):
            cell = hdr_row.cells[i]
            cell.width = Inches(widths[i])
            set_cell_background(cell, bg_header)
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            set_cell_borders(cell, top=HEX_PRIMARY, bottom=HEX_PRIMARY, sz=8)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        # Filas de datos
        for r_idx, row_values in enumerate(data):
            row = tbl.rows[r_idx + 1]
            bg_color = HEX_BG_ALT if (r_idx % 2 == 1) else "FFFFFF"
            for c_idx, val in enumerate(row_values):
                cell = row.cells[c_idx]
                cell.width = Inches(widths[c_idx])
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                set_cell_borders(cell, top=HEX_BORDER, bottom=HEX_BORDER, sz=4)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.1

                # Formato especial para primera columna (código)
                run = p.add_run(val)
                run.font.name = "Consolas" if c_idx == 0 else "Calibri"
                run.font.size = Pt(8.5 if c_idx == 0 else 9.0)
                if c_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = COLOR_PRIMARY
                elif c_idx == 1:
                    run.font.color.rgb = COLOR_SECONDARY
                else:
                    run.font.color.rgb = COLOR_TEXT

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(6)

    # ==============================================================================
    # 0. PORTADA EJECUTIVA ELEGANTE
    # ==============================================================================
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(20)
    p_pre.paragraph_format.space_after = Pt(4)
    run_pre = p_pre.add_run("SISTEMA DE GESTIÓN DE SEMICONDUCTORES Y MECATRÓNICA")
    run_pre.font.name = "Calibri"
    run_pre.font.size = Pt(11)
    run_pre.font.bold = True
    run_pre.font.color.rgb = COLOR_SECONDARY

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run("Arquitectura de Datos y Diagrama UML de Colecciones MongoDB")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("Modelado NoSQL de Documentos Raíz, Subdocumentos Embebidos (Embedding) y Relaciones Referenciadas (1:1 y 1:N) con django-mongodb-backend")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = COLOR_MUTED

    # Caja de metadatos del documento
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("Backend de Base de Datos:", " MongoDB Community / Atlas"), ("Framework Web:", " Django 6.1 (Python 3.12+)")],
        [("Conector Oficial:", " django-mongodb-backend"), ("Módulo / App:", " miapp (Catálogo y POS Mecatrónica)")]
    ]
    for row_idx, row_items in enumerate(meta_data):
        for col_idx, (k, v) in enumerate(row_items):
            cell = meta_table.cell(row_idx, col_idx)
            cell.width = Inches(3.35)
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
            set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz=4)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run_k = p.add_run(k)
            run_k.font.name = "Calibri"
            run_k.font.size = Pt(8.5)
            run_k.font.bold = True
            run_k.font.color.rgb = COLOR_PRIMARY
            run_v = p.add_run(v)
            run_v.font.name = "Calibri"
            run_v.font.size = Pt(8.5)
            run_v.font.color.rgb = COLOR_TEXT

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(8)
    p_space.paragraph_format.space_after = Pt(8)

    # ==============================================================================
    # 1. INTRODUCCIÓN Y PRINCIPIOS DE MODELADO EN MONGODB
    # ==============================================================================
    add_header_h1("Introducción y Principios de Modelado en MongoDB", "🏛️")

    add_paragraph(
        "El modelado de datos en MongoDB difiere fundamentalmente de los modelos relacionales (SQL). "
        "En lugar de normalizar datos a ultranza mediante docenas de tablas unidas por costosas operaciones "
        "de JOIN en tiempo de consulta, MongoDB y django-mongodb-backend permiten adoptar dos estrategias complementarias: "
        "la incrustación (Embedding) de subdocumentos y la referenciación (Referencing) mediante identificadores ObjectId."
    )

    add_paragraph(
        "En este sistema de comercio electrónico y punto de venta especializado en semiconductores y mecatrónica, "
        "la arquitectura fue diseñada meticulosamente considerando la atomicidad de las transacciones, la inmutabilidad "
        "histórica de las ventas y la contención de bloqueos durante picos de concurrencia."
    )

    add_bullet("1. Patrón de Documento Embebido (Embedded Documents): ", 
               "Se utiliza para datos fuertemente ligados a una entidad padre que se leen juntos y que no tienen sentido sin ella (ej. Direccion de un usuario, PrecioEscalonado de un componente y DetalleVenta de una orden). Permite lecturas atómicas en 1 sola operación BSON sin ningún JOIN.")
    
    add_bullet("2. Patrón de Referencia Normalizada (References): ", 
               "Se utiliza cuando las entidades crecen de forma independiente o son compartidas por múltiples consumidores (ej. Usuario y sus Órdenes, Categoría y Productos, Proveedores).")
    
    add_bullet("3. Patrón de Desacoplamiento de Inventario (Separated Collection): ", 
               "Se aísla la colección 'inventory' del catálogo 'products' en una relación 1 a 1 referenciada. Esto evita que escrituras masivas de bloqueo de stock (reservas, checkout, ventas en POS) generen bloqueos o invalidación de cachés en el catálogo público de productos.")

    add_callout(
        "Regla de Oro en MongoDB: 'Los datos que se acceden y consultan juntos, deben residir juntos'. "
        "Sin embargo, 'los datos con frecuencias de actualización radicalmente diferentes (lecturas de catálogo vs escrituras de stock) "
        "deben desacoplarse para garantizar escalabilidad'.",
        icon="⚖️",
        title="PRINCIPIO ARQUITECTÓNICO FUNDAMENTAL"
    )

    # ==============================================================================
    # 2. DIAGRAMA GRÁFICO UML DE CLASES Y COLECCIONES
    # ==============================================================================
    doc.add_page_break()
    add_header_h1("Diagrama Gráfico UML de Clases y Colecciones", "📊")

    add_paragraph(
        "A continuación se presenta el diagrama UML de clases y colecciones que modela la base de datos completa. "
        "El diagrama distingue visualmente entre colecciones raíz independientes de MongoDB, subdocumentos embebidos (EmbeddedModel) "
        "y colecciones auxiliares o desacopladas, señalando con exactitud la cardinalidad (1:1 y 1:N) y el tipo de relación."
    )

    # Inserción de imagen
    img_path = "uml_colecciones_mongodb.png"
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.8))

        p_caption = doc.add_paragraph()
        p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_caption.paragraph_format.space_before = Pt(2)
        p_caption.paragraph_format.space_after = Pt(12)
        run_cap = p_caption.add_run("Figura 1: Diagrama de Arquitectura UML - Colecciones, Subdocumentos Embebidos y Relaciones NoSQL.")
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(8.5)
        run_cap.font.italic = True
        run_cap.font.color.rgb = COLOR_MUTED

    add_header_h2("Guía de Notación y Convenciones del Diagrama")
    add_bullet("🟦 Rectángulos Azules («MongoDB Collection»): ", "Colecciones raíz de primer nivel en MongoDB con su propia clave primaria '_id: ObjectId' (db_table: users, categories, suppliers, products, orders, payments).")
    add_bullet("🟩 Rectángulos Verdes («Embedded Document»): ", "Subdocumentos incrustados directamente dentro del documento padre que no poseen colección propia (Direccion, PrecioEscalonado, DetalleVenta).")
    add_bullet("🟧 Rectángulo Ámbar («Inventory»): ", "Colección de almacén y stock físico desacoplada deliberadamente del catálogo general para alta concurrencia.")
    add_bullet("➔ Flecha Azul Sólida: ", "Relación por referencia (vía clave foránea / ObjectId). Indica que un documento almacena la referencia del otro.")
    add_bullet("◆--➔ Rombo Verde y Línea Punteada: ", "Composición de MongoDB (Incrustación / Embedding). El rombo se ubica en el documento contenedor y la flecha apunta al subdocumento o array embebido.")

    # ==============================================================================
    # 3. ANÁLISIS EXHAUSTIVO DE RELACIONES (1 A 1 Y 1 A MUCHOS)
    # ==============================================================================
    doc.add_page_break()
    add_header_h1("Análisis de Relaciones: Cardinalidad y Mecanismo NoSQL", "🔗")

    add_paragraph(
        "En un sistema híbrido Django + MongoDB, las relaciones pueden materializarse como campos referenciados "
        "(ObjectId / ForeignKey / OneToOneField) o como subdocumentos embebidos (EmbeddedModelField / EmbeddedModelArrayField). "
        "A continuación se detalla la justificación técnica de cada relación del sistema."
    )

    add_header_h2("1. Relaciones Uno a Uno (1:1)")

    add_bullet("A. auth.User ↔ Usuario (1:1 Referenciado - OneToOneField)", 
               "Extensión del modelo de autenticación estándar de Django. 'auth.User' almacena las credenciales de sesión, tokens y flags administrativos ('is_staff', 'is_superuser'). El modelo 'Usuario' actúa como perfil de dominio de negocio en la colección 'users'. Permite compatibilidad total con el Admin de Django mientras el perfil reside en una colección personalizada de MongoDB.")

    add_bullet("B. Usuario ↔ Direccion (1:1 Embebido doble - EmbeddedModelField)", 
               "El modelo 'Usuario' contiene dos subdocumentos embebidos independientes: 'direccion_envio' y 'direccion_facturacion'. "
               "Al guardarse en la colección 'users', no se crean registros en otra tabla. Cuando el usuario inicia sesión o carga su panel, sus direcciones se leen de forma instantánea sin requerir consultas cruzadas.")

    add_bullet("C. Producto ↔ Inventario (1:1 Referenciado - OneToOneField / ObjectId)", 
               "Cada 'Producto' del catálogo tiene exactamente un registro en la colección 'inventory' referenciado por 'producto_id'. "
               "¿Por qué no embeber el inventario dentro del producto? Porque en un e-commerce y punto de venta concurrido, miles de usuarios leen el catálogo (lectura pesada), mientras los cajeros y la pasarela bloquean y actualizan cantidades de stock (escritura pesada). Mantener 'inventory' separado previene la contención de escrituras y bloqueos a nivel de documento sobre los productos.")

    add_bullet("D. Orden ↔ Direccion (1:1 Embebido Instantánea - EmbeddedModelField)", 
               "La colección 'orders' incluye 'direccion_entrega' como un subdocumento embebido. Esto constituye una instantánea (snapshot): si el cliente cambia su dirección de residencia semanas después en su perfil, la orden previamente despachada mantiene congelada la dirección exacta de despacho de aquel momento histórico.")

    add_header_h2("2. Relaciones Uno a Muchos (1:N)")

    add_bullet("A. Categoria ↔ Categoria (1:N Jerarquía Autorreferencial - ForeignKey 'self')", 
               "Permite construir un árbol infinito de clasificación para semiconductores y mecatrónica (ejemplo: 'Semiconductores > Microcontroladores > 32-Bits > ARM Cortex-M4 > STM32'). Cada categoría hija apunta a su padre mediante 'parent_id'. El método 'ruta_jerarquica()' resuelve la ruta completa.")

    add_bullet("B. Categoria ↔ Producto (1:N Referenciado - ForeignKey / ObjectId)", 
               "Una categoría agrupa a muchos productos del catálogo, pero cada producto pertenece a una categoría principal. Al consultar productos por categoría, MongoDB indexa 'categoria_id' para filtrados ultra rápidos.")

    add_bullet("C. Proveedor ↔ Producto (1:N Referenciado - ForeignKey / ObjectId)", 
               "Un fabricante o proveedor (como Texas Instruments o STMicroelectronics) suministra cientos de componentes. La colección 'suppliers' guarda sus datos corporativos y líneas de especialidad, mientras los productos guardan la referencia 'proveedor_id'.")

    add_bullet("D. Producto ↔ PrecioEscalonado (1:N Embebido en Array - EmbeddedModelArrayField)", 
               "En el mercado de semiconductores, los componentes tienen precios por escala de volumen (Tiers: 1-9 uds a $120.00, 10-49 uds a $105.00, 50+ uds a $88.00). Estos rangos se almacenan como un Array embebido 'precios_escalonados' dentro del documento del producto. Cero JOINs: el punto de venta (POS) y el carrito calculan el precio aplicable en memoria instantáneamente.")

    add_bullet("E. Usuario ↔ Orden (1:N Referenciado - ForeignKey / ObjectId)", 
               "Un cliente puede generar cientos de órdenes a lo largo del tiempo. Las órdenes no se embeben dentro del usuario porque superarían el límite de 16 MB por documento BSON de MongoDB (antipatrón de arrays sin límite). Por ello, se normaliza: cada 'Orden' contiene la referencia 'usuario_id'.")

    add_bullet("F. Orden ↔ DetalleVenta (1:N Embebido en Array - EmbeddedModelArrayField)", 
               "Cada 'Orden' contiene el array embebido 'items' con subdocumentos 'DetalleVenta'. Cada elemento almacena una instantánea completa (producto_id, SKU, nombre, encapsulado, cantidad comprada, precio unitario aplicado y subtotal). Si el catálogo de productos cambia de precio o se descataloga en el futuro, las órdenes históricas permanecen intactas e inmutables.")

    add_bullet("G. Orden ↔ Pago (1:N Referenciado - ForeignKey / ObjectId)", 
               "Una orden puede recibir múltiples intentos de pago (un intento fallido de Stripe, una confirmación posterior por transferencia SPEI, o reembolsos parciales). Por tanto, la colección 'payments' guarda cada transacción vinculada a 'orden_id'.")

    # Tabla resumen de relaciones
    add_header_h2("Matriz Resumen de Relaciones y Cardinalidad")
    rel_headers = ["Entidad Origen", "Entidad Destino", "Cardinalidad", "Mecanismo MongoDB", "Propósito Arquitectónico"]
    rel_data = [
        ["auth.User", "Usuario", "1 : 1", "Referenciado (OneToOne)", "Separar credenciales de auth del perfil de negocio"],
        ["Usuario", "Direccion", "1 : 1 (doble)", "Embebido (EmbeddedField)", "Direcciones de envío y facturación sin JOINs"],
        ["Producto", "Inventario", "1 : 1", "Referenciado (OneToOne)", "Desacoplar escrituras de stock de lecturas de catálogo"],
        ["Orden", "Direccion", "1 : 1", "Embebido (EmbeddedField)", "Instantánea (snapshot) histórica de entrega"],
        ["Categoria", "Categoria", "1 : N", "Referenciado (FK self)", "Árbol jerárquico recursivo de componentes"],
        ["Categoria", "Producto", "1 : N", "Referenciado (ForeignKey)", "Clasificación de catálogo por categoría"],
        ["Proveedor", "Producto", "1 : N", "Referenciado (ForeignKey)", "Asociación de fabricante a productos"],
        ["Producto", "PrecioEscalonado", "1 : N", "Embebido (ArrayField)", "Precios por volumen (Tiers) dentro del producto"],
        ["Usuario", "Orden", "1 : N", "Referenciado (ForeignKey)", "Historial de pedidos sin límite de 16MB"],
        ["Orden", "DetalleVenta", "1 : N", "Embebido (ArrayField)", "Instantánea inmutable de productos comprados"],
        ["Orden", "Pago", "1 : N", "Referenciado (ForeignKey)", "Múltiples pasarelas, reintentos y conciliaciones"],
    ]
    add_table_fields(rel_headers, rel_data, [1.1, 1.1, 0.8, 1.5, 2.2])

    # ==============================================================================
    # 4. DICCIONARIO DETALLADO DE COLECCIONES Y SUBDOCUMENTOS
    # ==============================================================================
    doc.add_page_break()
    add_header_h1("Diccionario Detallado de Colecciones y Documentos", "📚")

    add_paragraph(
        "A continuación se desglosan en detalle cada una de las 7 colecciones raíz y los 3 subdocumentos embebidos "
        "implementados en el código de 'miapp/models.py', detallando sus campos, tipos BSON/Django, restricciones "
        "de negocio y ejemplos en formato JSON/BSON."
    )

    # ----------------------------------------------------
    # 4.1 COLECCIÓN: users (Usuario)
    # ----------------------------------------------------
    add_header_h2("1. Colección 'users' — Modelo: Usuario", badge="Colección Raíz")
    add_paragraph(
        "Almacena la información de perfiles, datos de contacto, rol en el sistema y las direcciones embebidas. "
        "Se relaciona 1:1 con la cuenta de autenticación de Django ('auth.User')."
    )
    user_fields = [
        ["_id", "ObjectId", "Obligatorio (PK)", "Identificador único nativo generado automáticamente por MongoDB."],
        ["user_id", "ObjectId", "Opcional (1:1)", "Clave foránea hacia 'auth.User' para sincronizar login y permisos."],
        ["username", "String", "Obligatorio (Unique)", "Identificador alfanumérico único para login y menciones."],
        ["email", "String", "Obligatorio (Unique)", "Correo electrónico verificado para notificaciones y facturación."],
        ["nombre_completo", "String", "Obligatorio (max 200)", "Nombre y apellidos completos del cliente o colaborador."],
        ["telefono", "String", "Opcional (max 30)", "Número de contacto telefónico para logística y paquetería."],
        ["rol", "String (Enum)", "Default: 'cliente'", "Roles: 'admin', 'cajero', 'cliente', 'distribuidor', 'ingeniero'."],
        ["direccion_envio", "Object (Embedded)", "Opcional", "Subdocumento embebido 'Direccion' para paquetería."],
        ["direccion_facturacion", "Object (Embedded)", "Opcional", "Subdocumento embebido 'Direccion' para emisión fiscal."],
        ["activo", "Boolean", "Default: True", "Flag de control para habilitar o suspender cuentas de usuario."],
        ["creado_en", "ISODate", "AutoNowAdd", "Marca de tiempo de registro del usuario en la plataforma."]
    ]
    add_table_fields(["Campo", "Tipo BSON", "Restricción", "Descripción y Regla de Negocio"], user_fields, [1.5, 1.2, 1.3, 2.7])
    
    add_json_box(
'''{
  "_id": ObjectId("65c1a8e9f1a2b3c4d5e6f701"),
  "username": "roberto_maker",
  "email": "roberto@mecatronica.org",
  "nombre_completo": "Roberto Gómez Bolaños",
  "telefono": "+52 55 9876 5432",
  "rol": "ingeniero",
  "direccion_envio": {
    "calle": "Av. Insurgentes Sur 1602, Piso 4",
    "colonia": "Crédito Constructor",
    "ciudad": "Benito Juárez",
    "estado": "Ciudad de México",
    "codigo_postal": "03940",
    "pais": "México",
    "referencias": "Junto al laboratorio de robótica"
  },
  "activo": true,
  "creado_en": ISODate("2026-03-15T10:30:00Z")
}''', caption="Ejemplo BSON: Documento en colección 'users'")

    # ----------------------------------------------------
    # 4.2 COLECCIÓN: categories (Categoria)
    # ----------------------------------------------------
    add_header_h2("2. Colección 'categories' — Modelo: Categoria", badge="Colección Raíz")
    add_paragraph(
        "Clasificación jerárquica en árbol recursivo para organizar componentes electrónicos, sensores y actuadores mecatrónicos."
    )
    cat_fields = [
        ["_id", "ObjectId", "Obligatorio (PK)", "Identificador único de la categoría."],
        ["nombre", "String", "Obligatorio (max 150)", "Nombre visible de la categoría (ej: 'Microcontroladores ARM')."],
        ["slug", "String", "Obligatorio (Unique)", "Identificador amigable para URLs indexables (ej: 'arm-cortex-m4')."],
        ["descripcion", "String", "Opcional (Text)", "Descripción técnica de la familia de componentes agrupados."],
        ["parent_id", "ObjectId", "Opcional (FK self)", "Referencia a la categoría padre superior para armar el árbol."],
        ["activa", "Boolean", "Default: True", "Visibilidad pública en el menú y filtros de búsqueda del catálogo."]
    ]
    add_table_fields(["Campo", "Tipo BSON", "Restricción", "Descripción y Regla de Negocio"], cat_fields, [1.5, 1.2, 1.3, 2.7])

    add_json_box(
'''{
  "_id": ObjectId("65c1a8e9f1a2b3c4d5e6f710"),
  "nombre": "STM32 Cortex-M4",
  "slug": "stm32-cortex-m4",
  "descripcion": "Microcontroladores de alto rendimiento con DSP y FPU",
  "parent_id": ObjectId("65c1a8e9f1a2b3c4d5e6f709"),
  "activa": true
}''', caption="Ejemplo BSON: Documento en colección 'categories' (Subcategoría jerárquica)")

    # ----------------------------------------------------
    # 4.3 COLECCIÓN: suppliers (Proveedor)
    # ----------------------------------------------------
    add_header_h2("3. Colección 'suppliers' — Modelo: Proveedor", badge="Colección Anexa")
    add_paragraph(
        "Directorio de fabricantes internacionales de semiconductores, distribuidores oficiales y marcas mecatrónicas."
    )
    supp_fields = [
        ["_id", "ObjectId", "Obligatorio (PK)", "Identificador único del proveedor o fabricante."],
        ["nombre", "String", "Obligatorio (max 200)", "Razón social / marca (ej: 'Texas Instruments', 'STMicroelectronics')."],
        ["codigo", "String", "Obligatorio (Unique)", "Siglas internacionales del fabricante (ej: 'TI', 'ST', 'NXP', 'MCHP')."],
        ["sitio_web", "String (URL)", "Opcional", "Enlace al portal oficial corporativo del fabricante."],
        ["email_contacto", "String", "Opcional", "Correo de atención para compras mayoristas y garantías."],
        ["telefono", "String", "Opcional", "Teléfono de contacto o soporte de fábrica."],
        ["pais", "String", "Default: 'México'", "País de origen o matriz del fabricante."],
        ["lineas_componentes", "String", "Opcional (Text)", "Líneas de especialidad: amplificadores, DSP, drivers de motor, etc."],
        ["activo", "Boolean", "Default: True", "Estado del proveedor para compras y reabastecimiento."],
        ["creado_en", "ISODate", "AutoNowAdd", "Fecha de alta del proveedor en el sistema."]
    ]
    add_table_fields(["Campo", "Tipo BSON", "Restricción", "Descripción y Regla de Negocio"], supp_fields, [1.5, 1.2, 1.3, 2.7])

    # ----------------------------------------------------
    # 4.4 COLECCIÓN: products (Producto)
    # ----------------------------------------------------
    add_header_h2("4. Colección 'products' — Modelo: Producto", badge="Colección Raíz")
    add_paragraph(
        "Catálogo integral de componentes con especificaciones técnicas flexibles en formato JSON nativo de MongoDB "
        "y array embebido de precios escalonados por volumen."
    )
    prod_fields = [
        ["_id", "ObjectId", "Obligatorio (PK)", "Identificador único de producto en la base de datos."],
        ["sku", "String", "Obligatorio (Unique)", "Número de Parte del Fabricante (MPN) (ej: 'STM32F401RET6')."],
        ["nombre", "String", "Obligatorio (max 250)", "Título comercial descriptivo del componente mecatrónico."],
        ["categoria_id", "ObjectId", "Obligatorio (FK)", "Referencia foránea a la colección 'categories'."],
        ["proveedor_id", "ObjectId", "Opcional (FK)", "Referencia foránea al fabricante en la colección 'suppliers'."],
        ["precio_base", "Decimal128", "Obligatorio (12, 2)", "Precio unitario para compras al por menor (1 a N unidades)."],
        ["descripcion", "String", "Opcional (Text)", "Resumen de características, arquitectura interna y periféricos."],
        ["datasheet_url", "String (URL)", "Opcional", "Enlace oficial a la hoja de datos en formato PDF."],
        ["encapsulado", "String", "Opcional (max 100)", "Package físico del integrado (ej: 'LQFP-64', 'SOIC-8', 'TO-220')."],
        ["voltaje_operacion", "String", "Opcional (max 100)", "Rango de alimentación eléctrica admitido (ej: '1.7V - 3.6V')."],
        ["pines", "Integer", "Opcional", "Número total de pines / leads del encapsulado."],
        ["especificaciones_tecnicas", "Object (JSON)", "Default: {}", "Atributos dinámicos (frecuencia, flash, RAM, torque, etc.)."],
        ["precios_escalonados", "Array[Object]", "Default: []", "Lista embebida de subdocumentos 'PrecioEscalonado' (Tiers)."],
        ["activo", "Boolean", "Default: True", "Flag de publicación activa en el catálogo público."],
        ["creado_en / actualizado_en", "ISODate", "AutoNow", "Auditoría de alta y última modificación técnica."]
    ]
    add_table_fields(["Campo", "Tipo BSON", "Restricción", "Descripción y Regla de Negocio"], prod_fields, [1.5, 1.2, 1.3, 2.7])

    add_json_box(
'''{
  "_id": ObjectId("65c1a8e9f1a2b3c4d5e6f720"),
  "sku": "STM32F401RET6",
  "nombre": "MCU ARM Cortex-M4 84MHz 512KB Flash",
  "categoria_id": ObjectId("65c1a8e9f1a2b3c4d5e6f710"),
  "proveedor_id": ObjectId("65c1a8e9f1a2b3c4d5e6f715"),
  "precio_base": NumberDecimal("125.50"),
  "encapsulado": "LQFP-64",
  "voltaje_operacion": "1.7V - 3.6V",
  "pines": 64,
  "especificaciones_tecnicas": {
    "frecuencia_max": "84 MHz",
    "memoria_flash": "512 KB",
    "sram": "96 KB",
    "adc_canales": 16,
    "temporizadores": 8
  },
  "precios_escalonados": [
    { "cantidad_minima": 10, "precio_unitario": NumberDecimal("112.00") },
    { "cantidad_minima": 50, "precio_unitario": NumberDecimal("98.50") },
    { "cantidad_minima": 250, "precio_unitario": NumberDecimal("84.00") }
  ],
  "activo": true
}''', caption="Ejemplo BSON: Documento 'products' con JSON flexible y array embebido de tiers")

    # ----------------------------------------------------
    # 4.5 COLECCIÓN: inventory (Inventario)
    # ----------------------------------------------------
    add_header_h2("5. Colección 'inventory' — Modelo: Inventario", badge="Colección Desacoplada 1:1")
    add_paragraph(
        "Control físico de almacén, ubicación en estanterías/gavetas y gestión atómica de existencias. "
        "Desacoplada del catálogo de productos para soportar alta concurrencia de ventas sin generar bloqueos en el catálogo."
    )
    inv_fields = [
        ["_id", "ObjectId", "Obligatorio (PK)", "Identificador único del registro de inventario."],
        ["producto_id", "ObjectId", "Obligatorio (1:1 Unique)", "Referencia única obligatoria hacia el producto correspondiente."],
        ["stock_disponible", "Integer", "Default: 0", "Existencias libres para venta inmediata en tienda física o web."],
        ["stock_reservado", "Integer", "Default: 0", "Unidades apartadas en carritos de checkout en proceso de pago."],
        ["ubicacion_anaquel", "String", "Opcional (max 100)", "Identificador físico de anaquel/gaveta (ej: 'Pasillo B, Gaveta 14-C')."],
        ["almacen", "String", "Default: 'Almacén Central'", "Nombre de la nave o sucursal donde reside la mercancía."],
        ["stock_minimo_alerta", "Integer", "Default: 10", "Umbral crítico para disparar órdenes automáticas de reorden."],
        ["actualizado_en", "ISODate", "AutoNow", "Marca de tiempo de la última operación de inventario."]
    ]
    add_table_fields(["Campo", "Tipo BSON", "Restricción", "Descripción y Regla de Negocio"], inv_fields, [1.5, 1.2, 1.3, 2.7], bg_header=HEX_AMBER)

    add_paragraph("Lógica de Negocio y Métodos Implementados:", bold_prefix="Métodos Atómicos de Inventario: ")
    add_bullet("• reservar(cantidad): ", "Resta de 'stock_disponible' y suma a 'stock_reservado' de forma atómica para evitar sobreventa durante el proceso de pago.")
    add_bullet("• liberar_reserva(cantidad): ", "Devuelve las unidades al stock disponible si el checkout expira o el usuario vacía su carrito.")
    add_bullet("• confirmar_venta(cantidad): ", "Liquida definitivamente las unidades reservadas al completarse exitosamente la orden de pago.")
    add_bullet("• descontar_venta_directa(cantidad): ", "Descuenta directamente del stock disponible para cobros en caja física (POS) en tiempo real.")

    # ----------------------------------------------------
    # 4.6 COLECCIÓN: orders (Orden)
    # ----------------------------------------------------
    add_header_h2("6. Colección 'orders' — Modelo: Orden", badge="Colección Raíz")
    add_paragraph(
        "Representa la transacción comercial consolidada. Incluye el subdocumento de dirección de entrega "
        "y el array embebido 'items' que almacena la instantánea inmutable de los productos comprados."
    )
    ord_fields = [
        ["_id", "ObjectId", "Obligatorio (PK)", "Identificador único de la orden."],
        ["folio", "String", "Obligatorio (Unique)", "Código único visible de venta (ej: 'ORD-2026-004821')."],
        ["usuario_id", "ObjectId", "Obligatorio (FK)", "Referencia al cliente en la colección 'users'."],
        ["estado", "String (Enum)", "Default: 'pendiente'", "Estados: 'pendiente', 'pagada', 'preparando', 'enviada', 'entregada', 'cancelada'."],
        ["items", "Array[Object]", "Obligatorio", "Array de subdocumentos 'DetalleVenta' congelados al momento de compra."],
        ["subtotal", "Decimal128", "Obligatorio", "Suma de los subtotales de todos los items comprados."],
        ["costo_envio", "Decimal128", "Default: 0.00", "Tarifa de logística o mensajería calculada."],
        ["total", "Decimal128", "Obligatorio", "Monto neto a pagar (subtotal + costo_envio - descuentos)."],
        ["direccion_entrega", "Object (Embedded)", "Opcional", "Instantánea de la dirección física donde se entregará el pedido."],
        ["notas", "String", "Opcional (Text)", "Instrucciones especiales de entrega o advertencias de manejo."],
        ["creado_en / actualizado_en", "ISODate", "AutoNow", "Auditoría de generación y cambio de estado del pedido."]
    ]
    add_table_fields(["Campo", "Tipo BSON", "Restricción", "Descripción y Regla de Negocio"], ord_fields, [1.5, 1.2, 1.3, 2.7])

    add_json_box(
'''{
  "_id": ObjectId("65c1a8e9f1a2b3c4d5e6f750"),
  "folio": "ORD-2026-008412",
  "usuario_id": ObjectId("65c1a8e9f1a2b3c4d5e6f701"),
  "estado": "pagada",
  "items": [
    {
      "producto_id": "65c1a8e9f1a2b3c4d5e6f720",
      "sku": "STM32F401RET6",
      "nombre": "MCU ARM Cortex-M4 84MHz",
      "encapsulado": "LQFP-64",
      "cantidad": 10,
      "precio_unitario": NumberDecimal("112.00"),
      "subtotal": NumberDecimal("1120.00")
    },
    {
      "producto_id": "65c1a8e9f1a2b3c4d5e6f725",
      "sku": "L298N-MOD",
      "nombre": "Módulo Driver Puente H Dual L298N",
      "encapsulado": "Módulo PCB",
      "cantidad": 2,
      "precio_unitario": NumberDecimal("65.00"),
      "subtotal": NumberDecimal("130.00")
    }
  ],
  "subtotal": NumberDecimal("1250.00"),
  "costo_envio": NumberDecimal("150.00"),
  "total": NumberDecimal("1400.00"),
  "direccion_entrega": {
    "calle": "Av. Insurgentes Sur 1602",
    "ciudad": "Benito Juárez",
    "estado": "Ciudad de México",
    "codigo_postal": "03940",
    "pais": "México"
  },
  "creado_en": ISODate("2026-03-20T14:45:00Z")
}''', caption="Ejemplo BSON: Documento 'orders' con items embebidos y totales consolidados")

    # ----------------------------------------------------
    # 4.7 COLECCIÓN: payments (Pago)
    # ----------------------------------------------------
    add_header_h2("7. Colección 'payments' — Modelo: Pago", badge="Colección Anexa")
    add_paragraph(
        "Bitácora y conciliación de pasarelas de pago externas (Stripe, PayPal, SPEI, Terminal bancaria POS)."
    )
    pago_fields = [
        ["_id", "ObjectId", "Obligatorio (PK)", "Identificador único de la transacción."],
        ["orden_id", "ObjectId", "Obligatorio (FK)", "Referencia a la orden en la colección 'orders'."],
        ["pasarela", "String (Enum)", "Obligatorio", "Pasarelas: 'stripe', 'paypal', 'transferencia', 'tarjeta'."],
        ["transaccion_id", "String", "Obligatorio (Unique)", "Token, Charge ID o referencia bancaria (ej: 'ch_3N4y...')."],
        ["monto", "Decimal128", "Obligatorio", "Monto exacto cobrado a través de la pasarela."],
        ["moneda", "String", "Default: 'MXN'", "Moneda de la transacción (MXN, USD, EUR)."],
        ["estado", "String (Enum)", "Default: 'pendiente'", "Estados: 'pendiente', 'aprobado', 'rechazado', 'reembolsado'."],
        ["datos_conciliacion", "Object (JSON)", "Default: {}", "Respuesta cruda en JSON devuelta por la API de la pasarela."],
        ["fecha_transaccion", "ISODate", "AutoNowAdd", "Marca de tiempo del intento o confirmación del pago."]
    ]
    add_table_fields(["Campo", "Tipo BSON", "Restricción", "Descripción y Regla de Negocio"], pago_fields, [1.5, 1.2, 1.3, 2.7])

    # ----------------------------------------------------
    # 4.8 SUBDOCUMENTOS EMBEBIDOS (EmbeddedModel)
    # ----------------------------------------------------
    add_header_h2("8. Subdocumentos Embebidos (EmbeddedModel de MongoDB)", badge="Subdocumentos")
    add_paragraph(
        "Los siguientes modelos heredan de 'EmbeddedModel' de 'django-mongodb-backend'. No poseen colección propia "
        "ni clave primaria '_id' independiente; se incrustan como objetos BSON o arrays dentro de sus documentos contenedores."
    )

    # Direccion
    add_paragraph("Subdocumento: Direccion", bold_prefix="A. ")
    add_paragraph("Almacena la estructura estándar de localización geográfica. Se embebe 1:1 en 'Usuario' (envío y facturación) y en 'Orden' (entrega).")
    dir_fields = [
        ["calle", "String", "Obligatorio (max 255)", "Nombre de la vialidad, número exterior e interior."],
        ["colonia", "String", "Opcional (max 150)", "Asentamiento o colonia de la demarcación."],
        ["ciudad", "String", "Obligatorio (max 100)", "Ciudad o municipio de residencia."],
        ["estado", "String", "Obligatorio (max 100)", "Entidad federativa o provincia."],
        ["codigo_postal", "String", "Obligatorio (max 20)", "Código postal oficial."],
        ["pais", "String", "Default: 'México'", "País de destino."],
        ["referencias", "String", "Opcional (max 255)", "Instrucciones de entrega para el mensajero."]
    ]
    add_table_fields(["Campo", "Tipo BSON", "Restricción", "Descripción"], dir_fields, [1.5, 1.2, 1.3, 2.7], bg_header=HEX_TEAL)

    # PrecioEscalonado
    add_paragraph("Subdocumento: PrecioEscalonado", bold_prefix="B. ")
    add_paragraph("Permite estructurar escalas de precios por volumen. Se embebe 1:N como array en 'Producto.precios_escalonados'.")
    tier_fields = [
        ["cantidad_minima", "PositiveInteger", "Obligatorio", "Umbral mínimo de piezas para activar la tarifa."],
        ["precio_unitario", "Decimal128", "Obligatorio", "Precio con descuento por unidad a partir de dicha cantidad."]
    ]
    add_table_fields(["Campo", "Tipo BSON", "Restricción", "Descripción"], tier_fields, [1.5, 1.2, 1.3, 2.7], bg_header=HEX_TEAL)

    # DetalleVenta
    add_paragraph("Subdocumento: DetalleVenta", bold_prefix="C. ")
    add_paragraph("Instantánea (snapshot) de compra congelada. Se embebe 1:N como array en 'Orden.items'.")
    det_fields = [
        ["producto_id", "String", "Obligatorio", "ID original del producto en el catálogo."],
        ["sku", "String", "Obligatorio", "Número de parte del fabricante congelado al comprar."],
        ["nombre", "String", "Obligatorio", "Nombre comercial del componente en la fecha de compra."],
        ["encapsulado", "String", "Opcional", "Package físico del componente vendido."],
        ["cantidad", "PositiveInteger", "Default: 1", "Número de piezas adquiridas en el renglón."],
        ["precio_unitario", "Decimal128", "Obligatorio", "Precio unitario pactado al momento de compra."],
        ["subtotal", "Decimal128", "Obligatorio", "Total del renglón (cantidad × precio_unitario)."]
    ]
    add_table_fields(["Campo", "Tipo BSON", "Restricción", "Descripción"], det_fields, [1.5, 1.2, 1.3, 2.7], bg_header=HEX_TEAL)

    # ==============================================================================
    # 5. BENEFICIOS ARQUITECTÓNICOS Y RENDIMIENTO
    # ==============================================================================
    doc.add_page_break()
    add_header_h1("Beneficios Arquitectónicos y Consideraciones de Rendimiento", "🚀")

    add_paragraph(
        "La arquitectura implementada combina las mejores prácticas del diseño de esquemas en MongoDB "
        "con la robustez del framework web Django 6.1:"
    )

    add_bullet("1. Cero JOINs en Consultas Críticas: ", 
               "Al consultar un producto para renderizar su ficha técnica o calcular su precio en el punto de venta, "
               "los precios escalonados (tiers) vienen incrustados en el mismo documento. Igualmente, al consultar una orden "
               "para imprimir el ticket o consultar el historial, todo el detalle de items se recupera en una sola lectura de disco.")

    add_bullet("2. Inmutabilidad y Auditoría Fiscal Garantizada: ", 
               "El uso de subdocumentos 'DetalleVenta' en lugar de claves foráneas directas al catálogo de productos garantiza "
               "que si un administrador altera el precio de un microcontrolador mañana, las órdenes cerradas en el pasado "
               "conserven inalterados sus montos originales, protegiendo la consistencia contable.")

    add_bullet("3. Máxima Escalabilidad y Baja Contención de Bloqueos: ", 
               "Al separar la colección 'inventory' del catálogo de productos, las operaciones concurrentes de cajeros y checkout "
               "no compiten por bloqueos de escritura con las miles de lecturas simultáneas del catálogo web.")

    add_bullet("4. Indexación Optimizada: ", 
               "Campos críticos como 'username', 'email', 'slug', 'codigo', 'sku', 'folio' y 'transaccion_id' cuentan con "
               "índices únicos en MongoDB, asegurando búsquedas en tiempo logarítmico O(log N) e integridad referencial.")

    output_filename = "Documentacion_UML_Colecciones_MongoDB.docx"
    doc.save(output_filename)
    print(f"Document successfully created: {output_filename}")

if __name__ == "__main__":
    generate_uml_docx()
